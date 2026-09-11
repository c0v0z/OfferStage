from pathlib import Path
from io import BytesIO


MAX_DOCUMENT_BYTES = 5 * 1024 * 1024


def _decode_text(data: bytes) -> str:
    return data.decode("utf-8-sig", errors="replace").strip()


def extract_text(data: bytes, filename: str) -> str:
    """Extract text from supported resume/JD formats with a small size guard."""
    if not data:
        return ""
    if len(data) > MAX_DOCUMENT_BYTES:
        raise ValueError("文件大小不能超过 5 MB。")

    suffix = Path(filename or "").suffix.lower()
    if suffix in {".txt", ".md", ".text"}:
        return _decode_text(data)
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader

            reader = PdfReader(BytesIO(data))
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
            return text.strip()
        except Exception as exc:
            raise ValueError(f"PDF 解析失败：{exc}") from exc
    if suffix == ".docx":
        try:
            from docx import Document

            document = Document(BytesIO(data))
            paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    values = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if values:
                        paragraphs.append(" | ".join(values))
            return "\n".join(paragraphs).strip()
        except Exception as exc:
            raise ValueError(f"DOCX 解析失败：{exc}") from exc
    raise ValueError("暂不支持该文件类型，请使用 TXT、PDF 或 DOCX。")
